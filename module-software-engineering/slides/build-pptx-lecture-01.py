"""Build the Lecture 01 deck as a .pptx in the style of 2026-09-seL4-summit.pptx.

Usage: python build-pptx-lecture-01.py pptx-template-blends.pptx OUTPUT.pptx

The template (pptx-template-blends.pptx, made by make-pptx-template.py; the
full seL4 deck also works) is opened so that its Blends master, layouts,
theme fonts (Microsoft Sans Serif titles, Tahoma body), square purple/green
bullets, banner, and footer are inherited unchanged. Its slides and the unused
Collins masters are dropped. Diagrams are drawn with native shapes and
connectors, tables with native tables, and the Marp speaker notes (HTML
comments) become PowerPoint notes.
"""
import re
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR, MSO_LINE
from pptx.oxml.ns import qn

TEMPLATE, OUT = sys.argv[1], sys.argv[2]
FOOTER = "Agentic Software Engineering — SE Module, Lecture 01"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_PURPLE = RGBColor(0x31, 0x00, 0x66)   # the Marp decks' heading purple; used for standout text only

# ---------------------------------------------------------------- template
prs = Presentation(TEMPLATE)


def theme_name(master):
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            return re.search(r'name="([^"]*)"', rel.target_part.blob.decode("utf8", "ignore")).group(1)


blends = next(m for m in prs.slide_masters if theme_name(m) == "Blends")

# drop every existing slide
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
# drop the masters the seL4 deck no longer uses (nine Collins masters)
mIdLst = prs.slide_masters._sldMasterIdLst
for mId in list(mIdLst):
    if prs.part.related_part(mId.rId) is not blends.part:
        prs.part.drop_rel(mId.rId)
        mIdLst.remove(mId)

L = {lay.name: lay for lay in blends.slide_layouts}

# ---------------------------------------------------------------- helpers
RICH = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def runs(p, text, size=None, bold=None, italic=None, color=None, font=None):
    """Add runs to paragraph p, honoring **bold**, *italic* and `code` markup."""
    for tok in RICH.split(text):
        if not tok:
            continue
        b, i, code = bold, italic, False
        if tok.startswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*"):
            tok, i = tok[1:-1], True
        elif tok.startswith("`"):
            tok, code = tok[1:-1], True
        r = p.add_run()
        r.text = tok
        f = r.font
        if size:
            f.size = Pt(size)
        if b is not None:
            f.bold = b
        if i is not None:
            f.italic = i
        if code:
            f.name = "Consolas"
        elif font:
            f.name = font
        if color is not None:
            f.color.rgb = color
    return p


def pPr(p):
    return p._p.get_or_add_pPr()


def no_bullet(p):
    pr = pPr(p)
    pr.set("marL", "0")
    pr.set("indent", "0")
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar"):
        for el in pr.findall(qn(tag)):
            pr.remove(el)
    etree.SubElement(pr, qn("a:buNone"))


def numbered(p, indent_in=0.4):
    pr = pPr(p)
    pr.set("marL", str(int(Inches(indent_in))))
    pr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar", "a:buFont", "a:buClr", "a:buSzPct"):
        for el in pr.findall(qn(tag)):
            pr.remove(el)
    etree.SubElement(pr, qn("a:buSzPct")).set("val", "100000")
    bf = etree.SubElement(pr, qn("a:buFont"))
    bf.set("typeface", "+mn-lt")
    an = etree.SubElement(pr, qn("a:buAutoNum"))
    an.set("type", "arabicPeriod")


def space_before(p, pts):
    pr = pPr(p)
    sb = pr.find(qn("a:spcBef"))
    if sb is None:
        sb = etree.Element(qn("a:spcBef"))
        pr.insert(0, sb)
    for el in list(sb):
        sb.remove(el)
    pts_el = etree.SubElement(sb, qn("a:spcPts"))
    pts_el.set("val", str(int(pts * 100)))


def add_footer(slide, layout, number=True):
    for ph in layout.placeholders:
        t = ph.placeholder_format.type
        if t == PP_PLACEHOLDER.FOOTER or (number and t == PP_PLACEHOLDER.SLIDE_NUMBER):
            slide.shapes.clone_placeholder(ph)
    for ph in slide.placeholders:
        t = ph.placeholder_format.type
        if t == PP_PLACEHOLDER.FOOTER:
            ph.text_frame.text = FOOTER
        elif t == PP_PLACEHOLDER.SLIDE_NUMBER:
            p = ph.text_frame.paragraphs[0]
            fld = etree.SubElement(p._p, qn("a:fld"))
            fld.set("id", "{6E0AA622-F4CE-604D-A669-CD3D12FC535C}")
            fld.set("type", "slidenum")
            etree.SubElement(fld, qn("a:rPr")).set("lang", "en-US")
            etree.SubElement(fld, qn("a:t")).text = "‹#›"


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def new_slide(layout_name, title=None, title_size=36, note=None, number=True):
    lay = L[layout_name]
    s = prs.slides.add_slide(lay)
    add_footer(s, lay, number=number)
    if title is not None and s.shapes.title is not None:
        tf = s.shapes.title.text_frame
        tf.text = title
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(title_size)
    notes(s, note)
    return s


def body(slide, items, size=22, sub_size=18):
    """Fill the body placeholder. items: (kind, text[, level]) with kind in
    'b' (bullet), 'n' (numbered), 'p' (prose, no bullet), 'gap' (spacer)."""
    ph = slide.placeholders[1]
    tf = ph.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        kind, text = it[0], it[1]
        level = it[2] if len(it) > 2 else 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        sz = size if level == 0 else sub_size
        if kind == "p":
            no_bullet(p)
            runs(p, text, size=sz)
            space_before(p, 10)
        elif kind == "n":
            numbered(p)
            runs(p, text, size=sz)
            space_before(p, 8)
        elif kind == "gap":
            no_bullet(p)
            runs(p, " ", size=8)
        else:
            runs(p, text, size=sz)
            space_before(p, 8 if level == 0 else 4)
    return ph


def set_fill(shape, kind):
    f = shape.fill
    if kind is None or kind == "none":
        f.background()
        return
    f.solid()
    if kind == "white":
        f.fore_color.rgb = WHITE
    elif kind == "yellow":
        f.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
    elif kind == "purple":
        f.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    elif kind == "lavender":
        f.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_5
        f.fore_color.brightness = 0.4
    elif kind == "lavender-strong":
        f.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_5
    elif kind == "gray":
        f.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_3
        f.fore_color.brightness = 0.35
    elif isinstance(kind, RGBColor):
        f.fore_color.rgb = kind


def set_line(shape, color=BLACK, width=1.0, dash=None, theme=None):
    ln = shape.line
    if color is None and theme is None:
        ln.fill.background()
        return
    if theme is not None:
        ln.color.theme_color = theme
    else:
        ln.color.rgb = color
    ln.width = Pt(width)
    if dash:
        ln.dash_style = dash


def add_shadow(shape):
    spPr = shape._element.spPr
    eff = spPr.find(qn("a:effectLst"))
    if eff is None:
        eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"))
    for k, v in dict(blurRad="63500", dist="35921", dir="2700000", algn="ctr", rotWithShape="0").items():
        sh.set(k, v)
    etree.SubElement(sh, qn("a:schemeClr")).set("val", "bg2")


def fill_text(shape, lines, size=13, bold_first=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
              color=None, italic=None, font="Tahoma", margins=0.06):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margins)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(lines, str):
        lines = lines.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs(p, line, size=size, bold=(True if (bold_first and i == 0) else None), color=color, italic=italic, font=font)
    return shape


STYLES = {
    # kind: (autoshape, fill, line color/theme, line width, text color, shadow)
    "artifact": (MSO_SHAPE.RECTANGLE, "lavender", BLACK, 1.0, None, False),
    "process": (MSO_SHAPE.RECTANGLE, "purple", None, 0, WHITE, False),
    "result": (MSO_SHAPE.RECTANGLE, "yellow", BLACK, 1.0, None, True),
    "intent": (MSO_SHAPE.ROUNDED_RECTANGLE, "white", "accent1", 1.5, None, False),
    "decision": (MSO_SHAPE.DIAMOND, "white", BLACK, 1.0, None, False),
    "note": (MSO_SHAPE.RECTANGLE, "lavender", BLACK, 1.0, None, True),
    "standout": (MSO_SHAPE.ROUNDED_RECTANGLE, "lavender-strong", None, 0, DARK_PURPLE, False),
    "plain": (MSO_SHAPE.RECTANGLE, "none", None, 0, None, False),
    "quote": (MSO_SHAPE.RECTANGLE, "white", BLACK, 1.0, None, False),
    "callout": (MSO_SHAPE.RECTANGLE, "yellow", None, 0, None, False),
}


def box(slide, x, y, w, h, text, kind="artifact", size=13, bold_first=True, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, italic=None, margins=0.06):
    shp_type, fill, line, lw, tcolor, shadow = STYLES[kind]
    shp = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shp, fill)
    if line is None:
        shp.line.fill.background()
    elif line == "accent1":
        set_line(shp, theme=MSO_THEME_COLOR.ACCENT_1, width=lw)
    else:
        set_line(shp, color=line, width=lw)
    if shadow:
        add_shadow(shp)
    else:
        shp.shadow.inherit = False
    if shp_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.18
    fill_text(shp, text, size=size, bold_first=bold_first, align=align, anchor=anchor, color=(tcolor or BLACK), italic=italic, margins=margins)
    return shp


def tbox(slide, x, y, w, h, text, size=16, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bold=None, italic=None,
         fill=None, color=None, font="Tahoma"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        set_fill(tb, fill)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text if isinstance(text, list) else text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs(p, line, size=size, bold=bold, italic=italic, color=color, font=font)
        if i:
            space_before(p, 6)
    return tb


def set_ends(conn, tail=True, head=False):
    ln = conn.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        el = ln.find(qn(tag))
        if el is None:
            el = etree.SubElement(ln, qn(tag))
        el.set("type", "triangle" if on else "none")
        el.set("w", "med")
        el.set("len", "med")


def arrow(slide, points, width=1.5, dashed=False, color=BLACK, head=True):
    """Polyline of straight connectors through points; arrowhead on the last."""
    segs = []
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(width)
        if dashed:
            c.line.dash_style = MSO_LINE.DASH
        set_ends(c, tail=False)
        segs.append(c)
    if head:
        set_ends(segs[-1], tail=True)
    return segs


def label(slide, x, y, w, text, size=11, h=0.3, italic=None):
    tb = tbox(slide, x, y, w, h, text, size=size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill="white", italic=italic)
    return tb


def table(slide, x, y, w, rows, col_widths, size=13, row_h=0.32):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_widths):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, cell_text in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            runs(tf.paragraphs[0], cell_text, size=size, bold=(True if i == 0 else None), font="Tahoma")
    return shp


def standout(title, sentence, note=None):
    s = new_slide("Title Only", title, note=note)
    box(s, 1.6, 2.6, 10.1, 2.2, sentence, kind="standout", size=28, bold_first=False)
    return s


# ---------------------------------------------------------------- slides
# 1 — title
s = new_slide("Title Slide", number=False)
t = s.shapes.title
t.left, t.top, t.width, t.height = Inches(0.6), Inches(1.6), Inches(12.1), Inches(1.9)
t.text_frame.text = "Specifications, Realizations,\vand Conformance"
for p in t.text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = Pt(40)
        r.font.bold = True
for ph in list(s.placeholders):
    if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
        ph._element.getparent().remove(ph._element)
for yy in (3.75, 4.5):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.1), Inches(yy), Inches(11.2), Inches(yy))
    ln.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    ln.line.width = Pt(2.25)
tbox(s, 2.1, 3.82, 9.1, 0.6, "Agentic Software Engineering — Software-Engineering Module, Lecture 01",
     size=22, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tbox(s, 2.1, 4.75, 9.1, 0.5, "Meeting 1 of 6 · runs part 1 of the note-set demo · launches **Exercise 1**",
     size=18, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 2 — the one idea
s = new_slide("Title and Content", "The one idea",
              note="0–4 min. No demo yet. ls the demo repository: one specification, a CLAUDE.md, a process/ folder, no notes.")
body(s, [
    ("p", "A specification and a realization are two artifacts. **Conformance** is a relation between them; **verification** is the activity that checks it."),
    ("gap", ""),
    ("p", "When the check fails, one side or the other is changed, the check is run again, and the decision — which side moved, and why — is recorded."),
], size=24)

# 3 — the example
s = new_slide("Title and Content", "The example: a set of study notes",
              note="0–4 min. Why the example is small: six rules on one screen; every concept can be pointed at in a file. Project 0 is the same thing at the student's scale.")
body(s, [
    ("p", "One markdown file per article you read. Two questions:"),
    ("b", "How should a note be formatted?"),
    ("b", "Who checks that it is?"),
    ("p", "The first answer is a document that is not a note: a **specification**, `note-format-spec.md`. The second answer is this lecture."),
    ("p", "At the start the repository holds the specification (version 0.1, draft), a short `CLAUDE.md`, and three process documents. No notes."),
], size=22)

# 4 — the format specification
s = new_slide("Title Only", "The format specification, version 0.1",
              note="4–14 min. Keep this on screen while defining the four terms. Note \"Version: 0.1 (draft)\".")
table(s, 1.0, 1.55, 11.3, [
    ["Rule", "Requirement (condensed)"],
    ["R1", "YAML front-matter block; the opening `---` is line 1"],
    ["R2", "exactly two fields: `title`, `created` (ISO-8601)"],
    ["R3", "the body begins with a level-1 heading"],
    ["R4", "exactly one level-1 heading"],
    ["R5", "no skipped heading levels"],
    ["R6", "`index.md` links every note exactly once; every link resolves"],
], [1.2, 10.1], size=16, row_h=0.46)
tbox(s, 1.0, 5.2, 11.3, 0.6, "RFC 2119 keywords; numbered rules, so that any report can cite one.", size=18)

# 5 — four terms
s = new_slide("Title and Content", "Four terms")
body(s, [
    ("b", "**Specification (S)** — states, above the level of the artifact it governs, what the developer intends; fixes what matters and leaves the rest free"),
    ("b", "**Realization (R)** — an artifact built to satisfy S; *implementation* when R is code; a note or a document can also be one"),
    ("b", "**Conformance** — the relation that holds when R satisfies every property stated in S; a yes-or-no question about the pair"),
    ("b", "**Verification** — the activity of *trying* to confirm that R conforms to S; the result is evidence, and we need to know what it is evidence of"),
    ("p", "One S, many R: every conformant note is a different realization of the same rules."),
], size=20)

# 6 — the arrangement (diagram: spec-realization)
s = new_slide("Title Only", "The arrangement")


def arrangement_row(y0, I, S, R, V, O, conf, size=12):
    bI = box(s, 0.55, y0 + 0.1, 1.75, 0.75, I, kind="intent", size=size, bold_first=False)
    bS = box(s, 2.7, y0, 2.4, 0.95, S, kind="artifact", size=size)
    bR = box(s, 7.0, y0, 2.4, 0.95, R, kind="artifact", size=size)
    bV = box(s, 4.9, y0 + 1.45, 2.0, 0.8, V, kind="process", size=size)
    bO = box(s, 8.6, y0 + 1.45, 3.1, 0.8, O, kind="result", size=size)
    yc = y0 + 0.475
    arrow(s, [(2.3, yc), (2.7, yc)])
    arrow(s, [(5.1, yc), (7.0, yc)], dashed=True)
    label(s, 5.2, yc - 0.32, 1.7, conf, size=10, h=0.64)
    arrow(s, [(4.2, y0 + 0.95), (4.2, y0 + 1.2), (5.4, y0 + 1.2), (5.4, y0 + 1.45)])
    arrow(s, [(7.5, y0 + 0.95), (7.5, y0 + 1.2), (6.4, y0 + 1.2), (6.4, y0 + 1.45)])
    arrow(s, [(6.9, y0 + 1.85), (8.6, y0 + 1.85)])


arrangement_row(1.5, "Developer intent",
                "Specification S\nstates the required properties; omits everything else",
                "Realization R\nbuilt to satisfy S",
                "Verification\nchecks whether conformance holds",
                "Result: holds, or findings that cite rule IDs",
                "conformance:\nR satisfies every property S states")
sep = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.55), Inches(4.05), Inches(12.6), Inches(4.05))
sep.line.color.theme_color = MSO_THEME_COLOR.ACCENT_3
sep.line.width = Pt(1)
sep.line.dash_style = MSO_LINE.DASH
tbox(s, 11.85, 1.45, 1.4, 0.4, "in general", size=11, italic=True, align=PP_ALIGN.RIGHT)
tbox(s, 11.85, 4.15, 1.4, 0.4, "in the example", size=11, italic=True, align=PP_ALIGN.RIGHT)
arrangement_row(4.3, "Intent: readable, indexable study notes",
                "`note-format-spec.md`\nrules R1–R6",
                "`notes/royce-1970-waterfall-paper.md`\n(later: `check_notes.py` too)",
                "a person, an agent, or `check_notes.py`",
                "'R2: created date is not ISO-8601'",
                "each note satisfies R1–R6")

# 7 — the loader and the process documents (table + diagram: process-documents)
s = new_slide("Title Only", "The loader, and the process documents",
              note="4–14 min. Rules are numbered AUD-1, RPT-2, VER-1, so that reports can cite process rules the way they cite R-numbers.")
tbox(s, 1.0, 1.32, 11.4, 0.6, "`CLAUDE.md` names one governing document and three process documents, and states two rules: read everything before acting; change nothing without approval.", size=16)
table(s, 1.0, 1.95, 11.4, [
    ["Document", "Says", "Loaded"],
    ["`process/spec-audit.md` (AUD)", "how a specification's quality is assessed", "when asked"],
    ["`process/reporting.md` (RPT)", "the forms a report takes", "always"],
    ["`process/verification.md` (VER)", "how conformance is checked, and by what", "always"],
], [3.9, 5.5, 2.0], size=13, row_h=0.3)
bL = box(s, 0.7, 4.75, 2.7, 0.9, "`CLAUDE.md` — the loader\nnames the documents; two rules", kind="process", size=12)
targets = [
    (3.6, "`note-format-spec.md`\nthe governing document"),
    (4.42, "`process/spec-audit.md`\nAUD — run when asked"),
    (5.24, "`process/reporting.md`\nRPT — always loaded"),
    (6.06, "`process/verification.md`\nVER — always loaded"),
]
for ty, txt in targets:
    box(s, 4.2, ty, 3.4, 0.62, txt, kind="artifact", size=12)
    arrow(s, [(3.4, 5.2), (4.2, ty + 0.31)], width=1.25)
arrow(s, [(7.6, 4.1), (8.1, 4.1), (8.1, 4.6), (7.6, 4.6)], dashed=True)
label(s, 8.25, 4.2, 1.3, "audited by", size=11)
arrow(s, [(7.6, 4.9), (8.1, 4.9), (8.1, 5.4), (7.6, 5.4)], dashed=True)
label(s, 8.25, 5.0, 1.5, "gap list, RPT-1", size=11)
arrow(s, [(7.6, 6.3), (8.1, 6.3), (8.1, 5.7), (7.6, 5.7)], dashed=True)
label(s, 8.25, 5.78, 2.0, "conformance report,\nRPT-2", size=11, h=0.44)
arrow(s, [(7.6, 3.7), (10.4, 3.7), (10.4, 6.5), (7.6, 6.5)], dashed=True)
label(s, 10.55, 4.85, 2.3, "realizations verified\nagainst it", size=11, h=0.5)

# 8 — five properties as an audit
s = new_slide("Title Only", "Evaluating a specification: five properties, as an audit", title_size=32,
              note="14–20 min. The properties are not advice; they are a procedure in a file, cited by number. Run when the developer asks, before any plan, after any amendment.")
table(s, 1.0, 1.5, 11.4, [
    ["Property", "Rule", "The test"],
    ["Unambiguous", "AUD-1", "could two careful readers decide a clause differently?"],
    ["Internally consistent", "AUD-2", "do two clauses that constrain the same thing state their relationship?"],
    ["Externally consistent", "AUD-3", "across documents: one vocabulary; nothing one document requires that another's artifact cannot satisfy"],
    ["Complete for its level", "AUD-4", "the walkthrough: perform each operation on paper; what did you invent?"],
    ["Traceable", "AUD-5", "can a report, a test, a later document cite this clause — still, after the next change?"],
], [2.6, 1.1, 7.7], size=14, row_h=0.5)
tbox(s, 1.0, 5.15, 11.4, 1.5, "AUD-6: search within a document, between documents, and between a document and executing it.\nAUD-7: the output is a numbered gap list (RPT-1); then stop and wait for rulings.", size=16)

# 9 — standout: the audit
standout("The audit", "Before anything realizes the specification, the specification is examined.",
         note="20–34 min begins. Demo Segment 2. Plan mode.")

# 10 — the audit prompt
s = new_slide("Title Only", "The audit prompt",
              note="Room writes its own list for two minutes before the agent's appears; compare. Seeded findings, for your eyes: the scope sentence against R6 (index.md cannot satisfy R1); R2 against R3 (no rule that title and H1 agree); the filename the walkthrough must invent.")
q = box(s, 1.0, 1.55, 11.4, 2.9,
        "Read `note-format-spec.md` and the process documents. Run the audit in `process/spec-audit.md` on the specification and report per RPT-1: number every finding, place it, quote the text, categorize it, and give a recommended resolution. Include the walkthrough that AUD-4 asks for: write one conformant note on paper, step by step, and say what you had to invent. State which of the three places in AUD-6 you searched. Then stop and wait for my rulings.",
        kind="quote", size=17, bold_first=False, align=PP_ALIGN.LEFT, margins=0.25)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.55), Inches(0.1), Inches(2.9))
set_fill(bar, "purple")
bar.line.fill.background()
tbox(s, 1.0, 4.75, 11.4, 1.0, "The prompt says which document and which audit. What the audit is, and what its report looks like, is in the files.", size=18)

# 11 — three findings, three rulings
s = new_slide("Title Only", "Three findings, three rulings",
              note="git diff HEAD~1 -- note-format-spec.md; git log --oneline. If a finding is missed: one nudge each, in the script.")
table(s, 0.8, 1.5, 11.8, [
    ["Finding", "Rule", "Ruling", "Amendment"],
    ["the scope sentence covers every markdown file; R6 names `index.md`, which cannot satisfy R1", "AUD-2", "`index.md` is not a note", "scope statement: R1–R5 govern `notes/`; R6 alone governs the index"],
    ["R2 requires `title`; R3 requires an H1; nothing says they agree", "AUD-2", "they are equal; the field is authoritative", "R3 extended"],
    ["R6 links to note files; no rule says what a note's filename is", "AUD-4", "derive it from the title", "new R7: the slug of the title"],
], [4.1, 0.9, 2.9, 3.9], size=13, row_h=0.55)
tbox(s, 0.8, 4.65, 11.8, 1.3, "Amendments proposed per RPT-3 — old text, new text, rationale, version, changelog — approved, applied: version 0.1 → 1.0.0. The specification moved before any realization existed.", size=17)

# 12 — when conformance fails (diagram: conformance-repair)
s = new_slide("Title Only", "When conformance fails: change S or change R",
              note="34–40 min. S can be missing a rule, or two rules can conflict, or a clause can fail to say what was meant. Two ways to find defects in S: realistic examples; analysis — which is what the audit is.")
bV = box(s, 0.7, 2.55, 2.5, 0.95, "Verify R against S\n(person, agent, or program)", kind="artifact", size=12)
bQ = box(s, 3.45, 2.3, 2.4, 1.45, "Conformance\nholds?", kind="decision", size=12, bold_first=False, margins=0.0)
bD = box(s, 6.3, 1.45, 3.4, 0.85, "Record the result (RPT-2):\nverifier, specification version, clauses checked", kind="result", size=12)
bW = box(s, 6.3, 3.15, 2.7, 1.0, "Decide which side moves\na person decides; the decision is recorded", kind="process", size=12)
bRR = box(s, 9.9, 2.5, 2.9, 0.95, "Repair R\ne.g. fix the date, the H1, the heading levels", kind="artifact", size=12)
bAS = box(s, 9.9, 3.9, 2.9, 1.1, "Amend S (RPT-3)\npropose → approve → edit → version bump → changelog", kind="artifact", size=12)
arrow(s, [(3.2, 3.025), (3.45, 3.025)])
arrow(s, [(4.65, 2.3), (4.65, 1.875), (6.3, 1.875)])
label(s, 4.75, 1.62, 0.6, "yes", size=11, h=0.28)
arrow(s, [(5.85, 3.025), (6.05, 3.025), (6.05, 3.65), (6.3, 3.65)])
label(s, 5.75, 3.12, 0.5, "no", size=11, h=0.28)
arrow(s, [(9.0, 3.45), (9.45, 3.45), (9.45, 2.975), (9.9, 2.975)])
label(s, 8.95, 2.68, 1.05, "R is wrong", size=11, h=0.28)
arrow(s, [(9.0, 3.85), (9.45, 3.85), (9.45, 4.45), (9.9, 4.45)])
label(s, 8.7, 4.5, 1.6, "S is wrong or silent", size=11, h=0.28)
arrow(s, [(12.8, 2.975), (13.05, 2.975), (13.05, 5.3), (1.95, 5.3), (1.95, 3.5)])
arrow(s, [(11.35, 5.0), (11.35, 5.3)], head=False)
tbox(s, 0.7, 5.5, 12.0, 1.4, "A failed check reports a fact about the pair (S, R). It does not say which side is wrong. A person decides; the decision is recorded — in the changelog when S moves, in the commit message when R moves.", size=17)

# 13 — standout: the verification gate
standout("The verification gate", "A note written carelessly, in another editor, with nothing checking it.",
         note="40–52 min. Demo Segment 3: mkdir notes; paste the careless note as notes/royce-1970-waterfall-paper.md (the filename follows R7); write index.md by hand; \"check and report per RPT-2; change nothing\".")

# 14 — a conformance report
s = new_slide("Title Only", "A conformance report (RPT-2)",
              note="The R3 finding exists only because of the audit. Against 0.1 a report would have accepted the heading, and it would have been right — for 0.1. That is why a report says which version it checked against.")
tbox(s, 1.0, 1.35, 11.4, 0.8, "**Verifier:** agent (model, date). **Specification:** `note-format-spec.md` 1.0.0. **Checked:** R1–R7. **Not checked:** none.", size=16)
table(s, 1.0, 2.15, 11.4, [
    ["In the file", "Rule", "Repair"],
    ["`created: Sept 11, 2026`", "R2 (since 0.1)", "`2026-09-11`"],
    ["H1 `The Waterfall Paper`; `title: Royce 1970 Waterfall Paper`", "R3 (amended an hour ago)", "heading rewritten; the field is authoritative"],
    ["a second H1, `# My take`", "R4", "demoted to `##`"],
    ["`####` directly after `##`", "R5", "demoted to `###`"],
], [5.2, 2.6, 3.6], size=14, row_h=0.46)
tbox(s, 1.0, 4.85, 11.4, 0.6, "**Verdict:** nonconformant. Repairs proposed; none applied.", size=18)

# 15 — report before repair
s = new_slide("Title and Content", "Report before repair, then the gate")
body(s, [
    ("b", "RPT-4: nothing changes until the report has been read and a ruling given"),
    ("b", "The ruling: the specification is right; the note is wrong"),
    ("b", "The repair; the re-check passes; the commit names the operation"),
    ("p", "The passing re-check is what \"done\" means: the task was to make the check pass, not to edit the file."),
    ("p", "In the audit, S moved with no R present. Here, R moved with S held still. Both end with conformance re-established."),
], size=22)

# 16 — three kinds of verifier
s = new_slide("Title Only", "Three kinds of verifier",
              note="52–64 min. Demo Segment 4: the agent writes check_notes.py (R1–R5, R7 per note; --all adds R6; exit 0/1/2; docstring states version and clauses). Run --all. Sabotage the date; exit 1; restore; exit 0.")
table(s, 1.0, 1.5, 11.4, [
    ["", "Human", "Agent", "Algorithmic"],
    ["Cost per check", "minutes of attention", "tokens", "milliseconds"],
    ["Same result on repeat", "not guaranteed", "not guaranteed", "yes"],
    ["Decides mechanical clauses", "yes", "yes", "yes"],
    ["Decides judgment clauses", "yes", "yes", "**no**"],
    ["Composable into a gate", "no", "with effort", "yes, via the exit code"],
], [3.3, 2.7, 2.7, 2.7], size=15, row_h=0.5)
tbox(s, 1.0, 4.75, 11.4, 1.0, "VER-1. Because every rule has an identifier, the three are comparable: all three cite R2–R5 on the same note.", size=18)

# 17 — verifiers (diagram)
s = new_slide("Title Only", "Verifiers")
bS = box(s, 2.4, 1.5, 3.3, 0.85, "Specification S\n`note-format-spec.md`", kind="artifact", size=13)
bR = box(s, 7.6, 1.5, 3.3, 0.85, "Realization R\na note", kind="artifact", size=13)
bH = box(s, 0.7, 3.45, 3.7, 1.35, "Human verifier\ndecides every clause, including judgment\nminutes per check; attention-dependent", kind="artifact", size=12)
bA = box(s, 4.8, 3.45, 3.7, 1.35, "Agent verifier\nmechanical and judgment clauses\ntokens per check; may vary run to run", kind="artifact", size=12)
bC = box(s, 8.9, 3.45, 3.7, 1.35, "Algorithmic verifier: `check_notes.py`\nmechanical clauses only\nfree, immediate, deterministic; exit 0 / 1 / 2", kind="artifact", size=12)
bF = box(s, 2.9, 5.9, 7.5, 0.75, "Reports that cite the same rule identifiers (RPT-2)", kind="process", size=14)
for sx, tx in ((3.0, 2.55), (4.05, 6.65), (5.1, 10.75)):
    arrow(s, [(sx, 2.35), (tx, 3.45)], width=1.0)
for sx, tx in ((8.2, 2.9), (9.25, 6.9), (10.3, 11.0)):
    arrow(s, [(sx, 2.35), (tx, 3.45)], width=1.0)
for cx, tx in ((2.55, 4.0), (6.65, 6.65), (10.75, 9.3)):
    arrow(s, [(cx, 4.8), (tx, 5.9)], width=1.0)

# 18 — mechanical clauses, judgment clauses
s = new_slide("Title and Content", "Mechanical clauses, judgment clauses —\vand the verifier as a realization", title_size=30)
body(s, [
    ("b", "**VER-2.** A clause is *mechanical* when a program can decide it from the file alone, *judgment* when it needs a reader. Every clause of 1.0.0 is mechanical. A program decides what it can; the rest is named, never assumed covered."),
    ("b", "`check_notes.py` is itself a realization of the specification. If it implements a rule wrongly, its results are wrong in a way its output does not show."),
    ("b", "Who verifies the verifier: read it against the rules; test it on inputs with known violations; compare it with another verifier."),
], size=20)

# 19 — six ways
s = new_slide("Title and Content", "Six ways a verification result can be invalid")
body(s, [
    ("n", "Checked against the wrong **version** of S — a 1.0.0 verifier on 2.0.0 notes: a false pass"),
    ("n", "The verifier **mis-implements** a clause; it is a realization too"),
    ("n", "Checked from **memory** rather than from the document (the loader's first rule)"),
    ("n", "**Clauses not decided**: a pass says nothing about clauses the verifier does not decide"),
    ("n", "**Attention**: the human is the verifier for everything no process checks"),
    ("n", "A **defective S**: before the amendments, no `index.md` could conform"),
    ("p", "Verification is not proof. It is evidence within a scope — which verifier, which version, which clauses — and RPT-2 requires the scope to be stated."),
], size=18)

# 20 — what the process documents did today
s = new_slide("Title Only", "What the process documents did today", note="64–70 min.")
table(s, 1.0, 1.5, 11.4, [
    ["Document", "Used", "For"],
    ["AUD", "once, when asked, before anything was built", "the gap list"],
    ["RPT", "throughout", "a gap list (RPT-1), two conformance reports (RPT-2), an amendment proposal (RPT-3), report-before-repair (RPT-4)"],
    ["VER", "throughout", "three kinds of verifier (VER-1); mechanical and judgment clauses (VER-2)"],
], [1.6, 3.9, 5.9], size=14, row_h=0.5)
tbox(s, 1.0, 4.3, 11.4, 1.3, "Not yet present: operations on the set, a statement of what the set is *for*, the rule that conformance is an invariant while the set changes. Lecture 02.", size=18)

# 21 — exercise 1 and project 0
s = new_slide("Title and Content", "Exercise 1, and Project 0", note="70–72 min.")
body(s, [
    ("p", "**Exercise 1** — run the audit and verification segments yourself; write one note carelessly; check it with all three kinds of verifier; report with RPT-2 scope lines; explain one disagreement; say which side should have moved and how each verifier could have been wrong. Due before Lecture 03."),
    ("gap", ""),
    ("p", "**Project 0** — the same structure at your scale: your PKB specification is this specification grown up; its checklist is R1–R7 grown up; the stretch-goal validator is `check_notes.py` grown up."),
], size=20)

# 22 — questions
s = new_slide("Title and Content", "Questions to think about")
body(s, [
    ("n", "Name a property that specification 1.0.0 leaves unconstrained. Should it stay that way? If not, write the rule and say which kind of verifier decides it."),
    ("n", "For the R3 finding, defend moving S instead of R. What would it cost?"),
    ("n", "A compiler's type checker is an algorithmic verifier. Which failure modes apply, and who verifies that verifier?"),
], size=20)

# 23 — before next meeting
s = new_slide("Title and Content", "Before next meeting")
body(s, [
    ("b", "Read Royce (1970) and Meyer (1992); see `reading-list.md`. Meyer's precondition and postcondition are the form the next lecture gives to an operation."),
    ("b", "Read the Lecture 02 handout: the draft concept of operations and the three process documents Lecture 02 adds."),
    ("b", "Exercise 1 is due before Lecture 03; Project 0 continues."),
    ("p", "**Next meeting:** what the set is for; operations with contracts; the specification as an invariant while the set changes."),
], size=20)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))

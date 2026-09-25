"""Strip a deck down to a template: no slides, only the Blends master.

Usage: python make-pptx-template.py SOURCE.pptx TEMPLATE.pptx

Every slide is removed and every slide master except the one whose theme is
named "Blends" is dropped, together with the layouts, themes, and media that
only they referenced. The result opens in PowerPoint as an empty presentation
and is the template build-pptx-lecture-*.py scripts start from.
"""
import re
import sys
from pptx import Presentation

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)


def theme_name(master):
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            return re.search(r'name="([^"]*)"', rel.target_part.blob.decode("utf8", "ignore")).group(1)


keep = next(m for m in prs.slide_masters if theme_name(m) == "Blends")
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
mIdLst = prs.slide_masters._sldMasterIdLst
for mId in list(mIdLst):
    if prs.part.related_part(mId.rId) is not keep.part:
        prs.part.drop_rel(mId.rId)
        mIdLst.remove(mId)
prs.save(dst)
print(f"{dst}: masters={len(prs.slide_masters)} layouts={[l.name for l in keep.slide_layouts]} slides={len(prs.slides)}")
